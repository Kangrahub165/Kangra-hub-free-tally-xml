import io
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Optional, List, Dict, Any, Tuple, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.transactions.model import CanonicalStatement, TransactionItem
from app.transactions.snapshot import FinalConversionSnapshot
from app.core.exceptions import XMLGenerationException

EXCEL_COLUMNS = [
    "VOUCHER NO.",
    "VOUCHER TYPE",
    "VOUCHER DATE",
    "VOUCHER NARRATION",
    "LEDGER NAME",
    "BANK LEDGER AS TALLY",
    "INST. NUMBER",
    "INST. DATE",
    "DEBIT AMOUNT",
    "CREDIT AMOUNT",
    "BALANCE"
]

def sanitize_filename_part(text: str) -> str:
    """Removes invalid filename characters."""
    return re.sub(r'[^A-Za-z0-9_\-]', '', text)

def format_excel_date_label(d: Optional[date]) -> str:
    """Formats date as DD-Mon-YY e.g. 01-Jul-25 for filename and text labels."""
    if not d:
        return ""
    return d.strftime("%d-%b-%y")

class TallyExcelGenerator:
    """
    Generates high-precision Excel (.xlsx) workbooks directly from FinalConversionSnapshot.
    Strictly conforms to PRD architectural corrections:
    - Does NOT depend on XML generation
    - Consumes the single source of truth: FinalConversionSnapshot
    - 11 exact columns in exact order
    - Single worksheet named 'Tally Data'
    - Header row frozen (A2) with auto-filters enabled
    - Real numeric cells for DEBIT AMOUNT, CREDIT AMOUNT, BALANCE (accounting format #,##0.00)
    - Real date cells for VOUCHER DATE, INST. DATE (display format DD-MMM-YY)
    - Instrument numbers stored strictly as text cells to preserve leading zeros
    - No cell merging
    - Column auto-sizing
    """

    def __init__(self, default_bank_ledger: str = "Bank Account", default_cash_ledger: str = "Cash"):
        self.default_bank_ledger = default_bank_ledger
        self.default_cash_ledger = default_cash_ledger

    def generate_excel_bytes(
        self,
        data: Union[FinalConversionSnapshot, CanonicalStatement],
        bank_ledger_name: Optional[str] = None,
        cash_ledger_name: Optional[str] = None
    ) -> bytes:
        """Generates the .xlsx file content as bytes."""
        wb = self.generate_workbook(data, bank_ledger_name, cash_ledger_name)
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream.getvalue()

    def generate_workbook(
        self,
        data: Union[FinalConversionSnapshot, CanonicalStatement],
        bank_ledger_name: Optional[str] = None,
        cash_ledger_name: Optional[str] = None
    ) -> openpyxl.Workbook:
        """Constructs and styles the openpyxl Workbook directly from FinalConversionSnapshot."""
        # Normalize to FinalConversionSnapshot
        if isinstance(data, FinalConversionSnapshot):
            snapshot = data
        elif isinstance(data, CanonicalStatement):
            b_name = bank_ledger_name or data.bank_ledger_name or self.default_bank_ledger
            c_name = cash_ledger_name or data.cash_ledger_name or self.default_cash_ledger
            snapshot = FinalConversionSnapshot.create_from_statement(data, bank_ledger_name=b_name, cash_ledger_name=c_name)
        else:
            raise XMLGenerationException("Invalid data provided to TallyExcelGenerator.")

        if not snapshot.transactions:
            raise XMLGenerationException("No transactions available in snapshot to generate Excel export.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tally Data"

        # Freeze the header row
        ws.freeze_panes = "A2"

        # Style definitions
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
        thin_border_side = Side(border_style="thin", color="CBD5E1")
        border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        regular_font = Font(name="Calibri", size=10)
        align_left = Alignment(horizontal="left", vertical="center")
        align_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")

        # 1. Write Header Row
        ws.row_dimensions[1].height = 28
        for col_idx, col_name in enumerate(EXCEL_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border

        # 2. Write Data Rows
        # Sequential voucher numbers 1, 2, 3...
        for seq_idx, tx in enumerate(snapshot.transactions, start=1):
            row_idx = seq_idx + 1
            ws.row_dimensions[row_idx].height = 22

            # Voucher Type from snapshot
            vch_type = tx.voucher_type or "Payment"

            # Party Ledger from snapshot
            party_ledger = (tx.ledger_name or "Suspense").strip()

            # Bank Ledger As Tally from snapshot
            bank_ledger = snapshot.bank_ledger_name.strip()

            # Instrument / Cheque number (strictly text to preserve leading zeros)
            inst_no = tx.instrument_number or tx.cheque_number or ""
            inst_no_str = str(inst_no).strip() if inst_no else None

            # Instrument Date
            inst_date = tx.instrument_date

            # Amounts: strictly numeric, standard decimal accounting precision
            debit_val: Optional[float] = float(tx.debit) if (tx.debit and tx.debit > Decimal("0.00")) else None
            credit_val: Optional[float] = float(tx.credit) if (tx.credit and tx.credit > Decimal("0.00")) else None
            balance_val: Optional[float] = float(tx.balance) if tx.balance is not None else None

            # 1. VOUCHER NO. (Sequential 1, 2, 3...)
            c1 = ws.cell(row=row_idx, column=1, value=seq_idx)
            c1.font = regular_font
            c1.alignment = align_center
            c1.border = border

            # 2. VOUCHER TYPE
            c2 = ws.cell(row=row_idx, column=2, value=vch_type)
            c2.font = regular_font
            c2.alignment = align_center
            c2.border = border

            # 3. VOUCHER DATE (Actual Excel Date, display DD-MMM-YY)
            c3 = ws.cell(row=row_idx, column=3)
            if tx.date:
                c3.value = tx.date
                c3.number_format = "DD-mmm-YY"
            else:
                c3.value = None
            c3.font = regular_font
            c3.alignment = align_center
            c3.border = border

            # 4. VOUCHER NARRATION (Complete unabbreviated narration with cell wrap enabled)
            full_narr = tx.narration or getattr(tx, "original_narration", "") or ""
            c4 = ws.cell(row=row_idx, column=4, value=full_narr)
            c4.font = regular_font
            c4.alignment = align_left_wrap
            c4.border = border

            # 5. LEDGER NAME
            c5 = ws.cell(row=row_idx, column=5, value=party_ledger)
            c5.font = regular_font
            c5.alignment = align_left
            c5.border = border

            # 6. BANK LEDGER AS TALLY (Exact configured ledger)
            c6 = ws.cell(row=row_idx, column=6, value=bank_ledger)
            c6.font = regular_font
            c6.alignment = align_left
            c6.border = border

            # 7. INST. NUMBER (Preserve leading zeros strictly as text)
            c7 = ws.cell(row=row_idx, column=7)
            if inst_no_str:
                c7.value = inst_no_str
                c7.data_type = "s"
                c7.number_format = "@"
            else:
                c7.value = None
            c7.font = regular_font
            c7.alignment = align_center
            c7.border = border

            # 8. INST. DATE (Actual Excel Date, display DD-MMM-YY)
            c8 = ws.cell(row=row_idx, column=8)
            if inst_date:
                c8.value = inst_date
                c8.number_format = "DD-mmm-YY"
            else:
                c8.value = None
            c8.font = regular_font
            c8.alignment = align_center
            c8.border = border

            # 9. DEBIT AMOUNT (Numeric Excel cell)
            c9 = ws.cell(row=row_idx, column=9)
            if debit_val is not None:
                c9.value = debit_val
                c9.number_format = "#,##0.00"
            else:
                c9.value = None
            c9.font = regular_font
            c9.alignment = align_right
            c9.border = border

            # 10. CREDIT AMOUNT (Numeric Excel cell)
            c10 = ws.cell(row=row_idx, column=10)
            if credit_val is not None:
                c10.value = credit_val
                c10.number_format = "#,##0.00"
            else:
                c10.value = None
            c10.font = regular_font
            c10.alignment = align_right
            c10.border = border

            # 11. BALANCE (Numeric Excel cell)
            c11 = ws.cell(row=row_idx, column=11)
            if balance_val is not None:
                c11.value = balance_val
                c11.number_format = "#,##0.00"
            else:
                c11.value = None
            c11.font = regular_font
            c11.alignment = align_right
            c11.border = border

        total_rows = len(snapshot.transactions) + 1
        # Enable Excel Auto-Filter on header row
        ws.auto_filter.ref = f"A1:K{total_rows}"

        # Column Auto-Sizing
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = cell.value
                if val is None:
                    continue
                if isinstance(val, date):
                    cell_len = 9  # e.g. "17-Jul-25"
                elif isinstance(val, (int, float, Decimal)):
                    cell_len = len(f"{val:,.2f}")
                else:
                    cell_len = len(str(val))
                if cell_len > max_len:
                    max_len = cell_len
            
            col_width = min(max(max_len + 4, 12), 65)
            ws.column_dimensions[col_letter].width = col_width

        return wb

    def generate_filename(
        self,
        data: Union[FinalConversionSnapshot, CanonicalStatement],
        bank_name: Optional[str] = None
    ) -> str:
        """
        Generates standard sanitized filename:
        {Bank}_Statement_{From}_to_{To}.xlsx
        Example: SBI_Statement_01-Jul-25_to_31-Jul-25.xlsx
        """
        raw_bank = bank_name or getattr(data, "bank_name", None) or getattr(data, "bank", None) or "Bank"
        first_token = raw_bank.split()[0]
        clean_bank = sanitize_filename_part(first_token) or "Bank"

        from_str = format_excel_date_label(data.statement_from) if data.statement_from else "Statement"
        to_str = format_excel_date_label(data.statement_to) if data.statement_to else "Export"

        return f"{clean_bank}_Statement_{from_str}_to_{to_str}.xlsx"
